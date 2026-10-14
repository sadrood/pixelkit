
import sys, json, os
infile = sys.argv[1]
ext = os.path.splitext(infile)[1].lower()
max_chars = int(sys.argv[2]) if len(sys.argv) > 2 else 50000

text = ""
try:
    if ext in ('.txt', '.md', '.csv', '.log', '.json', '.xml', '.html', '.htm', '.css', '.js', '.ts', '.py', '.java', '.c', '.cpp', '.h', '.sh', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.conf'):
        with open(infile, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    elif ext == '.pdf':
        import fitz
        doc = fitz.open(infile)
        pages = []
        for page in doc:
            t = page.get_text()
            if t.strip():
                pages.append(t.strip())
        doc.close()
        text = '\n\n--- Page Break ---\n\n'.join(pages)
    elif ext in ('.docx', '.doc'):
        from docx import Document
        doc = Document(infile)
        paras = []
        for p in doc.paragraphs:
            if p.text.strip():
                # 检测标题样式
                if p.style and p.style.name and 'Heading' in p.style.name:
                    level = p.style.name.split()[-1]
                    prefix = '#' * (int(level) if level.isdigit() else 1) + ' '
                    paras.append(prefix + p.text.strip())
                else:
                    paras.append(p.text.strip())
        text = '\n\n'.join(paras)
    elif ext in ('.xlsx', '.xls'):
        import openpyxl
        wb = openpyxl.load_workbook(infile, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f'## Sheet: {sheet_name}\n')
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(c) if c is not None else '' for c in row]
                lines.append(' | '.join(row_vals))
        text = '\n'.join(lines)
    else:
        # 尝试当作文本文件读取
        with open(infile, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
except Exception as e:
    print(json.dumps({"success": False, "error": f"读取失败: {str(e)}"}))
    sys.exit(0)

total_len = len(text)
text = text[:max_chars]

print(json.dumps({"success": True, "text": text, "totalLength": total_len, "truncated": total_len > max_chars}))
