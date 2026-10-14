
import sys, json, os
from markdownify import markdownify as md
from docx import Document

infile = sys.argv[1]
output = sys.argv[2]
ext = os.path.splitext(infile)[1].lower()

text = ""

if ext in ('.docx', '.doc'):
    doc = Document(infile)
    md_lines = []
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            level = int(para.style.name.split()[-1]) if para.style.name.split()[-1].isdigit() else 1
            md_lines.append('#' * level + ' ' + para.text)
        else:
            md_lines.append(para.text)
    text = '\n\n'.join(md_lines)

elif ext in ('.html', '.htm'):
    with open(infile, 'r', encoding='utf-8') as f:
        html = f.read()
    text = md(html, heading_style='ATX')

elif ext == '.xlsx':
    import openpyxl
    wb = openpyxl.load_workbook(infile, data_only=True)
    md_lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        md_lines.append(f'## {sheet_name}\n')
        rows_data = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows_data):
            row_vals = [str(c) if c is not None else '' for c in row]
            md_lines.append('| ' + ' | '.join(row_vals) + ' |')
            if i == 0:
                md_lines.append('|' + '|'.join(['---' for _ in row_vals]) + '|')
        md_lines.append('')
    text = '\n'.join(md_lines)

else:
    with open(infile, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()

with open(output, 'w', encoding='utf-8') as f:
    f.write(text)

print(json.dumps({"success": True}))
